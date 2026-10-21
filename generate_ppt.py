import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt

def add_slide(prs, layout, title, bullet_points):
    slide = prs.slides.add_slide(layout)
    shapes = slide.shapes
    
    title_shape = shapes.title
    title_shape.text = title
    
    body_shape = shapes.placeholders[1]
    tf = body_shape.text_frame
    
    for point in bullet_points:
        p = tf.add_paragraph()
        p.text = point
        p.level = 0
        
    return slide

prs = Presentation()

# Slide 1: Title Slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "CodeForge"
subtitle.text = "A Lightweight Command-Line Text Editor in C++\n\nProject Overview & Architecture"

# Slide 2: Introduction
add_slide(prs, prs.slide_layouts[1], "Introduction", [
    "What is CodeForge? A terminal-based text editor built from scratch in C++.",
    "Runs in 'raw mode' to directly process user input without waiting for Enter.",
    "Uses standard POSIX terminal controls (termios, ioctl).",
    "Lightweight, fast, minimalistic, and dependency-free."
])

# Slide 3: Project Architecture
add_slide(prs, prs.slide_layouts[1], "Project Architecture", [
    "main.cpp: Entry point, parses command-line arguments and initializes the editor.",
    "Editor Class: Core component handling terminal I/O, UI rendering, and user logic.",
    "Buffer Class: Data structure holding the lines of text in memory.",
    "Cursor Class: Manages the row/col coordinates of the text cursor on screen."
])

# Slide 4: Terminal Management (Editor)
add_slide(prs, prs.slide_layouts[1], "Terminal Management", [
    "Raw Mode: Disables echo and canonical mode using tcgetattr & tcsetattr.",
    "Escape Sequences: Interprets VT100 ANSI escape sequences for terminal control.",
    "Responsive Window: Uses ioctl(TIOCGWINSZ) to detect terminal window size.",
    "Double Buffering Concept: Renders text to an in-memory string before writing to stdout."
])

# Slide 5: Buffer Management
add_slide(prs, prs.slide_layouts[1], "Buffer Management", [
    "In-Memory Representation: Uses std::vector<std::string> for storing lines.",
    "Text Operations: Fast insert/delete of characters and lines.",
    "Line Splitting/Joining: Splits lines on Enter, joins them on Backspace/Delete.",
    "Scrolling Logic: Maintains vertical and horizontal offsets when text exceeds terminal size."
])

# Slide 6: Key Features & Controls
add_slide(prs, prs.slide_layouts[1], "Key Features & Controls", [
    "Navigation: Arrow keys, Page Up/Down, Home, End.",
    "Saving: Ctrl+S saves the current buffer to disk.",
    "Quitting: Ctrl+Q safely exits the editor.",
    "Status Bar: Displays current filename, total lines, and cursor coordinates.",
    "Message Bar: Shows transient notifications (e.g., save status)."
])

# Slide 7: Potential Future Enhancements
add_slide(prs, prs.slide_layouts[1], "Future Enhancements", [
    "Syntax Highlighting for C/C++ and other languages.",
    "Search and Replace functionality.",
    "Support for multiple open buffers or split panes.",
    "Undo/Redo history stack.",
    "Configuration files (.codeforgerc) for custom keybindings."
])

prs.save('CodeForge_Presentation.pptx')
print("Presentation generated successfully at CodeForge_Presentation.pptx")
